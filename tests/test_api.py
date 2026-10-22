import io
import json
import zipfile
import pytest
import pymupdf as fitz
from docx import Document
from fastapi.testclient import TestClient
from backend.app import app
from backend import db

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(db,'DATA',tmp_path/'data')
    with TestClient(app) as c:yield c

def create(c,text='求 $\\int x e^x dx$。'):
    r=c.post('/api/questions',json={'text':text,'number':'1','subject':'数学一','confirmed':True})
    assert r.status_code==200;return r.json()['id']

def test_question_history_notes_and_backup(client):
    a=create(client);b=create(client,'另一道题')
    r=client.post('/api/math/integrate',json={'expression':'x*exp(x)','question_id':a});assert r.status_code==200,r.text
    result=r.json();assert result['verified'];hid=result['history_id']
    assert len(client.get('/api/math/history',params={'question_id':a}).json())==1
    assert client.get('/api/math/history',params={'question_id':b}).json()==[]
    assert client.post(f'/api/questions/{b}/calculations/{hid}',json={}).status_code==409
    assert client.post(f'/api/questions/{a}/calculations/{hid}',json={}).status_code==200
    assert client.post(f'/api/questions/{a}/calculations/{hid}',json={}).status_code==200
    body={'answer':'(x-1)e^x+C','notes':'先分部积分','draft':'[[{"x":1,"y":2}]]','wrong':True,'status':'已完成'}
    assert client.put(f'/api/questions/{a}/attempt',json=body).status_code==200
    reopened=client.get(f'/api/questions/{a}').json();assert reopened['attempt']['notes']==body['notes'];assert len(reopened['calculations'])==1
    assert reopened['calculations'][0]['result']['steps']==result['steps']
    assert len(client.get('/api/questions',params={'wrong':True}).json())==1
    stats=client.get('/api/stats').json();assert stats['questions']==2 and stats['completed']==1 and stats['wrong']==1
    archive=zipfile.ZipFile(io.BytesIO(client.get('/api/backup').content));assert 'mathmaster.db' in archive.namelist()

def pdf_bytes():
    d=fitz.open();page=d.new_page();page.insert_text((50,90),'1. Integrate x*exp(x).',fontsize=14);page.insert_text((50,170),'2. Differentiate sin(x^2).',fontsize=14)
    data=d.tobytes();d.close();return data

def test_pdf_import_review_split_merge_crop(client):
    source=pdf_bytes()
    r=client.post('/api/papers/import',data={'subject':'数学二','title':'本地测试试卷'},files={'file':('exam.pdf',source,'application/pdf')})
    assert r.status_code==200,r.text;paper=r.json();assert paper['count']==2
    questions=client.get('/api/questions',params={'paper_id':paper['id']}).json();assert len(questions)==2
    a,b=questions;assert not a['confirmed'] and a['images']
    assert client.get('/api/files/'+a['images'][0]).status_code==200
    assert client.get('/api/files/'+a['source_filename']).content==source
    original_pixels=client.get('/api/files/'+a['images'][0]).content
    crop=client.post(f"/api/questions/{a['id']}/crop",json={'image':a['images'][0],'top':0,'bottom':0.9});assert crop.status_code==200
    assert client.get('/api/files/'+a['images'][0]).content==original_pixels # exact original kept
    assert client.post(f"/api/questions/{a['id']}/merge",json={'other_id':b['id']}).status_code==200
    merged=client.get('/api/questions/'+a['id']).json();assert 'Differentiate' in merged['text']
    assert client.post(f"/api/questions/{a['id']}/split",json={'position':merged['text'].index('2.')}).status_code==200
    assert len(client.get('/api/questions',params={'paper_id':paper['id']}).json())==2
    backup=zipfile.ZipFile(io.BytesIO(client.get('/api/backup').content));assert 'uploads/'+a['source_filename'] in backup.namelist()

def test_docx_import(client):
    d=Document();d.add_paragraph('1. Calculate x^2.');d.add_paragraph('2. Calculate determinant.');out=io.BytesIO();d.save(out)
    r=client.post('/api/papers/import',data={'subject':'数学三'},files={'file':('exam.docx',out.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert r.status_code==200,r.text;assert r.json()['count']==2
    qs=client.get('/api/questions',params={'paper_id':r.json()['id']}).json();assert qs[0]['source_filename'].endswith('.docx');assert not qs[0]['confirmed']

def test_scanned_pdf_keeps_images(client):
    from PIL import Image,ImageDraw
    image=Image.new('RGB',(700,900),'white');ImageDraw.Draw(image).text((50,50),'1. Integrate x^2.\n\n2. Find the rank.',fill='black',font_size=24)
    buf=io.BytesIO();image.save(buf,format='PNG');pdf=fitz.open();page=pdf.new_page();page.insert_image(page.rect,stream=buf.getvalue())
    r=client.post('/api/papers/import',files={'file':('scan.pdf',pdf.tobytes(),'application/pdf')})
    assert r.status_code==200,r.text
    qs=client.get('/api/questions',params={'paper_id':r.json()['id']}).json()
    assert qs and all(q['images'] and not q['confirmed'] for q in qs)
    from backend.importer import tesseract_binary
    if tesseract_binary():assert len(qs)>=2 and all(q['text'].strip() for q in qs)
    assert all(client.get('/api/files/'+q['images'][0]).status_code==200 for q in qs)

def test_input_security_and_unknown_question(client):
    assert client.post('/api/math/integrate',json={'expression':"__import__('os')"}).status_code==422
    assert client.post('/api/math/integrate',json={'expression':'x','question_id':'missing'}).status_code==404
    assert client.get('/api/files/../../requirements.txt').status_code==404
    assert client.post('/api/questions',json={'text':'hello'},headers={'Origin':'https://example.com'}).status_code==403
    assert client.post('/api/papers/import',files={'file':('evil.exe',b'bad')}).status_code==422

def test_offline_sequential_calculations(client,monkeypatch):
    import socket
    def no_network(*args,**kwargs):raise AssertionError('External network is forbidden')
    monkeypatch.setattr(socket,'create_connection',no_network)
    for op,body,expected in [('integrate',{'expression':'x^2','lower':'0','upper':'1'},'1/3'),('differentiate',{'expression':'sin(x^2)'},'2*x*cos(x**2)'),('matrix',{'matrix':[['1','2'],['3','4']],'action':'determinant'},'-2')]:
        r=client.post('/api/math/'+op,json=body);assert r.status_code==200,r.text;assert r.json()['exact_result']==expected;assert r.json()['verified']
    assert len(client.get('/api/math/history',params={'workspace':True}).json())==3


def test_concurrent_note_write_conflict(client):
    q=create(client)
    first={'notes':'窗口 A 的笔记','expected_updated_at':None}
    r=client.put('/api/questions/'+q+'/attempt',json=first);assert r.status_code==200
    conflict=client.put('/api/questions/'+q+'/attempt',json={'notes':'窗口 B 的旧笔记','expected_updated_at':None})
    assert conflict.status_code==409
    assert client.get('/api/questions/'+q).json()['attempt']['notes']=='窗口 A 的笔记'
    update=client.put('/api/questions/'+q+'/attempt',json={'notes':'窗口 A 的新笔记','expected_updated_at':r.json()['updated_at']})
    assert update.status_code==200

def test_implicit_api_history_and_validation(client):
    qid=create(client)
    body={'expression':'x^2+y^2=1','mode':'implicit','dependent_variable':'y',
          'point':{'x':'3/5','y':'4/5'},'question_id':qid}
    r=client.post('/api/math/differentiate',json=body)
    assert r.status_code==200,r.text
    result=r.json();assert result['verified'] and result['exact_result']=='-3/4'
    h=client.get('/api/math/history',params={'question_id':qid}).json()[0]
    assert h['request']['mode']=='implicit' and h['request']['point']==body['point']
    assert h['result']['steps']==result['steps']
    assert client.post(f"/api/questions/{qid}/calculations/{result['history_id']}",json={}).status_code==200
    for extra in [{'order':2},{'variables':['x']},{'dependent_variable':'x'},{'mode':'unknown'},{'point':{'x':'0'}}]:
        assert client.post('/api/math/differentiate',json={**body,**extra}).status_code==422
    assert client.post('/api/math/differentiate',json={'expression':'x=y'}).status_code==422
    assert len(client.get('/api/math/history',params={'question_id':qid}).json())==1
